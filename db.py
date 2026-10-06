import hashlib
import json
import os
import sqlite3
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

DB_PATH = Path(__file__).with_name("toonscripture.db")

# Publishable credentials are intentionally safe to ship in app code.
# RLS + Supabase Auth protect the actual data.
SUPABASE_URL = os.getenv(
    "SUPABASE_URL",
    "https://wuiedibvewelsrrupeyh.supabase.co",
)
SUPABASE_PUBLISHABLE_KEY = os.getenv(
    "SUPABASE_PUBLISHABLE_KEY",
    "sb_publishable_AcgklpwumKyD20fOPixgZg_yiFTG6kn",
)


def using_supabase():
    return bool(SUPABASE_URL and SUPABASE_PUBLISHABLE_KEY)


def _cloud_client():
    from supabase import create_client

    client = create_client(SUPABASE_URL, SUPABASE_PUBLISHABLE_KEY)

    access_token = os.getenv("TOONSCRIPTURE_ACCESS_TOKEN", "").strip()
    refresh_token = os.getenv("TOONSCRIPTURE_REFRESH_TOKEN", "").strip()

    if access_token and refresh_token:
        client.auth.set_session(access_token, refresh_token)

    return client


def auth_sign_in(email, password):
    client = _cloud_client()
    response = client.auth.sign_in_with_password(
        {"email": email.strip(), "password": password}
    )
    session = response.session
    user = response.user

    if not session or not user:
        raise RuntimeError("Sign-in did not return a valid session.")

    # Confirm this account is approved for ToonScripture.
    client.auth.set_session(session.access_token, session.refresh_token)
    result = (
        client.table("admin_users")
        .select("user_id")
        .eq("user_id", user.id)
        .execute()
    )

    if not result.data:
        client.auth.sign_out()
        raise PermissionError("This account is not approved for ToonScripture.")

    return {
        "access_token": session.access_token,
        "refresh_token": session.refresh_token,
        "email": user.email,
        "user_id": user.id,
    }


def auth_restore(access_token, refresh_token):
    client = _cloud_client()
    response = client.auth.set_session(access_token, refresh_token)
    session = response.session

    if not session:
        raise RuntimeError("Could not restore your session.")

    user = response.user or client.auth.get_user().user

    return {
        "access_token": session.access_token,
        "refresh_token": session.refresh_token,
        "email": getattr(user, "email", None),
        "user_id": getattr(user, "id", None),
    }


def auth_sign_out(access_token=None, refresh_token=None):
    try:
        client = _cloud_client()
        if access_token and refresh_token:
            client.auth.set_session(access_token, refresh_token)
        client.auth.sign_out()
    except Exception:
        pass


def _client():
    """
    Return a Supabase client carrying the current Streamlit user's session.
    The tokens stay in that browser session instead of global process state.
    """
    client = _cloud_client()

    try:
        import streamlit as st

        access_token = st.session_state.get("auth_access_token")
        refresh_token = st.session_state.get("auth_refresh_token")

        if access_token and refresh_token:
            restored = client.auth.set_session(
                access_token,
                refresh_token,
            )

            if restored.session:
                st.session_state.auth_access_token = (
                    restored.session.access_token
                )
                st.session_state.auth_refresh_token = (
                    restored.session.refresh_token
                )
    except Exception:
        pass

    return client


# -------------------------------------------------------------------
# LOCAL SQLITE FALLBACK
# -------------------------------------------------------------------

def _local_connect():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def _local_init_db():
    con = _local_connect()
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS projects(
            id TEXT PRIMARY KEY,
            story_name TEXT NOT NULL,
            bible_reference TEXT,
            target_minutes REAL DEFAULT 6,
            format TEXT DEFAULT 'long_form',
            status TEXT DEFAULT 'idea',
            treatment_json TEXT,
            script_json TEXT,
            critique_json TEXT,
            character_bible_json TEXT,
            scenes_json TEXT,
            package_json TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS catalog_people(
            id TEXT PRIMARY KEY,
            source_sheet TEXT,
            source_position INTEGER,
            testament TEXT,
            book TEXT,
            name TEXT NOT NULL,
            entity_type TEXT,
            named_status TEXT,
            bible_references TEXT,
            story_role TEXT,
            priority TEXT,
            cinematic_score REAL,
            youtube_hook TEXT,
            reference_asset_need TEXT,
            primary_environment TEXT,
            canon TEXT,
            playlist_series TEXT,
            notes TEXT,
            source_url TEXT,
            production_status TEXT DEFAULT 'not_started',
            completed_at TEXT,
            imported_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS catalog_series(
            id TEXT PRIMARY KEY,
            rank INTEGER,
            series_name TEXT UNIQUE NOT NULL,
            series_type TEXT,
            core_concept TEXT,
            best_starter_episodes TEXT,
            youtube_reason TEXT,
            recommended_scene_count TEXT,
            recommended_runtime TEXT,
            priority TEXT,
            status TEXT,
            source_notes TEXT,
            imported_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS project_catalog_links(
            project_id TEXT PRIMARY KEY,
            catalog_person_id TEXT,
            catalog_series_id TEXT,
            episode_title TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS references_library(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id TEXT NOT NULL,
            reference_type TEXT NOT NULL,
            name TEXT NOT NULL,
            status TEXT DEFAULT 'not_approved',
            master_prompt TEXT,
            identity_lock TEXT,
            negative_lock TEXT,
            image_path TEXT,
            metadata_json TEXT,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(project_id, reference_type, name)
        );
        """
    )
    con.commit()
    con.close()


def init_db():
    if using_supabase():
        return
    _local_init_db()


def _local_row(row):
    return dict(row) if row else None


# -------------------------------------------------------------------
# PROJECTS
# -------------------------------------------------------------------

def list_projects():
    if using_supabase():
        result = (
            _client()
            .table("projects")
            .select("*")
            .order("created_at", desc=True)
            .execute()
        )
        return result.data or []

    con = _local_connect()
    rows = [
        dict(r)
        for r in con.execute(
            "SELECT * FROM projects ORDER BY created_at DESC"
        )
    ]
    con.close()
    return rows


def get_project(project_id):
    if not project_id:
        return None

    if using_supabase():
        result = (
            _client()
            .table("projects")
            .select("*")
            .eq("id", project_id)
            .limit(1)
            .execute()
        )
        return result.data[0] if result.data else None

    con = _local_connect()
    row = con.execute(
        "SELECT * FROM projects WHERE id=?",
        (project_id,),
    ).fetchone()
    con.close()
    return _local_row(row)


def create_project(
    story_name,
    bible_reference="",
    target_minutes=6.0,
    fmt="long_form",
):
    project_id = str(uuid.uuid4())[:8]

    payload = {
        "id": project_id,
        "story_name": story_name.strip(),
        "bible_reference": bible_reference.strip(),
        "target_minutes": float(target_minutes),
        "format": fmt,
        "status": "idea",
    }

    if using_supabase():
        _client().table("projects").insert(payload).execute()
        return project_id

    con = _local_connect()
    con.execute(
        """
        INSERT INTO projects(
            id, story_name, bible_reference, target_minutes, format
        )
        VALUES(?,?,?,?,?)
        """,
        (
            project_id,
            payload["story_name"],
            payload["bible_reference"],
            payload["target_minutes"],
            payload["format"],
        ),
    )
    con.commit()
    con.close()
    return project_id


def save_json(project_id, field, value, status=None):
    allowed = {
        "treatment_json",
        "script_json",
        "critique_json",
        "character_bible_json",
        "scenes_json",
        "package_json",
    }

    if field not in allowed:
        raise ValueError(field)

    if using_supabase():
        payload = {
            field: value,
        }
        if status:
            payload["status"] = status

        (
            _client()
            .table("projects")
            .update(payload)
            .eq("id", project_id)
            .execute()
        )
        return

    con = _local_connect()
    if status:
        con.execute(
            f"""
            UPDATE projects
            SET {field}=?, status=?, updated_at=CURRENT_TIMESTAMP
            WHERE id=?
            """,
            (json.dumps(value), status, project_id),
        )
    else:
        con.execute(
            f"""
            UPDATE projects
            SET {field}=?, updated_at=CURRENT_TIMESTAMP
            WHERE id=?
            """,
            (json.dumps(value), project_id),
        )

    con.commit()
    con.close()


def load_json(project, field, default=None):
    if not project:
        return default

    raw = project.get(field)

    if raw in (None, ""):
        return default

    if isinstance(raw, (dict, list)):
        return raw

    try:
        return json.loads(raw)
    except Exception:
        return default


# -------------------------------------------------------------------
# EPISODE PERFORMANCE
# -------------------------------------------------------------------

def save_episode_metrics(
    project_id,
    views=None,
    impressions=None,
    ctr=None,
    average_view_duration_seconds=None,
    average_percentage_viewed=None,
    subscribers_gained=None,
    notes=None,
):
    if not using_supabase():
        return None

    payload = {
        "project_id": project_id,
        "views": views,
        "impressions": impressions,
        "ctr": ctr,
        "average_view_duration_seconds": average_view_duration_seconds,
        "average_percentage_viewed": average_percentage_viewed,
        "subscribers_gained": subscribers_gained,
        "notes": notes,
    }

    result = (
        _client()
        .table("episode_metrics")
        .insert(payload)
        .execute()
    )
    return result.data[0] if result.data else payload


def list_episode_metrics(project_id):
    if not using_supabase():
        return []

    result = (
        _client()
        .table("episode_metrics")
        .select("*")
        .eq("project_id", project_id)
        .order("recorded_at", desc=True)
        .execute()
    )
    return result.data or []


def record_ai_usage(
    project_name,
    section_name,
    model_name,
    outcome,
    http_status=None,
):
    if not using_supabase():
        return None

    payload = {
        "project_name": project_name,
        "section_name": section_name,
        "model_name": model_name,
        "outcome": outcome,
        "http_status": http_status,
    }

    try:
        _client().table("ai_usage_events").insert(payload).execute()
    except Exception:
        pass

    return payload


def ai_usage_today():
    if not using_supabase():
        return []

    from datetime import datetime, timezone

    start = datetime.now(timezone.utc).replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    ).isoformat()

    result = (
        _client()
        .table("ai_usage_events")
        .select("*")
        .gte("created_at", start)
        .order("created_at", desc=True)
        .execute()
    )

    rows = result.data or []
    summary = {}

    for row in rows:
        name = row.get("project_name") or "Unknown"
        item = summary.setdefault(
            name,
            {
                "project_name": name,
                "requests": 0,
                "successes": 0,
                "errors": 0,
                "last_status": None,
                "last_section": None,
                "last_model": None,
                "last_http_status": None,
                "last_used_at": None,
            },
        )

        item["requests"] += 1

        if row.get("outcome") == "success":
            item["successes"] += 1
        else:
            item["errors"] += 1

        if item["last_used_at"] is None:
            item["last_status"] = row.get("outcome")
            item["last_section"] = row.get("section_name")
            item["last_model"] = row.get("model_name")
            item["last_http_status"] = row.get("http_status")
            item["last_used_at"] = row.get("created_at")

    return list(summary.values())


# -------------------------------------------------------------------
# GENERIC PROJECT ASSETS
# -------------------------------------------------------------------

def save_project_asset(
    project_id,
    asset_type,
    name,
    file_bytes,
    filename,
    content_type="application/octet-stream",
    metadata=None,
):
    if not using_supabase():
        raise RuntimeError("Cloud storage is required for project assets.")

    client = _client()

    ext = ""
    if "." in filename:
        ext = "." + filename.rsplit(".", 1)[1].lower()

    safe_name = _slug(name)
    path = (
        f"{_slug(project_id)}/"
        f"{_slug(asset_type)}/"
        f"{safe_name}-{uuid.uuid4().hex[:10]}{ext}"
    )

    (
        client.storage
        .from_("toonscripture-assets")
        .upload(
            path=path,
            file=file_bytes,
            file_options={
                "content-type": content_type,
                "upsert": "false",
            },
        )
    )

    payload = {
        "project_id": project_id,
        "asset_type": asset_type,
        "name": name,
        "storage_path": path,
        "metadata_json": metadata or {},
    }

    result = (
        client.table("project_assets")
        .insert(payload)
        .execute()
    )

    return result.data[0] if result.data else payload


def list_project_assets(project_id, asset_type=None):
    if not using_supabase():
        return []

    query = (
        _client()
        .table("project_assets")
        .select("*")
        .eq("project_id", project_id)
    )

    if asset_type:
        query = query.eq("asset_type", asset_type)

    result = query.order("created_at", desc=True).execute()
    return result.data or []


def signed_asset_url(storage_path, expires_in=3600):
    if not using_supabase() or not storage_path:
        return None

    signed = (
        _client().storage
        .from_("toonscripture-assets")
        .create_signed_url(storage_path, expires_in)
    )

    if isinstance(signed, dict):
        return signed.get("signedURL") or signed.get("signed_url")

    return None


# -------------------------------------------------------------------
# REFERENCE ASSETS
# -------------------------------------------------------------------

def _slug(value):
    value = str(value or "").strip().lower()
    safe = []
    for ch in value:
        if ch.isalnum():
            safe.append(ch)
        elif ch in (" ", "-", "_"):
            safe.append("-")
    slug = "".join(safe)
    while "--" in slug:
        slug = slug.replace("--", "-")
    return slug.strip("-") or "reference"


def save_reference_asset(
    project_id,
    reference_type,
    name,
    file_bytes,
    filename,
    content_type="application/octet-stream",
    master_prompt=None,
    identity_lock=None,
    negative_lock=None,
):
    if not using_supabase():
        raise RuntimeError("Reference uploads require the cloud database.")

    client = _client()

    ext = ""
    if "." in filename:
        ext = "." + filename.rsplit(".", 1)[1].lower()

    path = (
        f"{_slug(project_id)}/"
        f"{_slug(reference_type)}/"
        f"{_slug(name)}-{uuid.uuid4().hex[:10]}{ext}"
    )

    (
        client.storage
        .from_("toonscripture-assets")
        .upload(
            path=path,
            file=file_bytes,
            file_options={
                "content-type": content_type,
                "upsert": "false",
            },
        )
    )

    payload = {
        "project_id": project_id,
        "reference_type": reference_type,
        "name": name,
        "status": "approved",
        "master_prompt": master_prompt,
        "identity_lock": identity_lock,
        "negative_lock": negative_lock,
        "image_path": path,
        "metadata_json": {
            "original_filename": filename,
            "content_type": content_type,
        },
    }

    (
        client.table("references_library")
        .upsert(
            payload,
            on_conflict="project_id,reference_type,name",
        )
        .execute()
    )

    return payload


def get_reference_asset(project_id, reference_type, name):
    if not using_supabase():
        return None

    result = (
        _client()
        .table("references_library")
        .select("*")
        .eq("project_id", project_id)
        .eq("reference_type", reference_type)
        .eq("name", name)
        .limit(1)
        .execute()
    )

    if not result.data:
        return None

    item = result.data[0]

    if item.get("image_path"):
        try:
            signed = (
                _client().storage
                .from_("toonscripture-assets")
                .create_signed_url(
                    item["image_path"],
                    3600,
                )
            )
            if isinstance(signed, dict):
                item["signed_url"] = (
                    signed.get("signedURL")
                    or signed.get("signed_url")
                )
        except Exception:
            pass

    return item


def list_reference_assets(project_id):
    if not using_supabase():
        return []

    result = (
        _client()
        .table("references_library")
        .select("*")
        .eq("project_id", project_id)
        .order("reference_type")
        .execute()
    )
    return result.data or []


# -------------------------------------------------------------------
# CATALOGUE
# -------------------------------------------------------------------

def catalog_summary():
    if using_supabase():
        people = (
            _client()
            .table("catalog_people")
            .select("id,entity_type,priority")
            .execute()
            .data
            or []
        )
        series = (
            _client()
            .table("catalog_series")
            .select("id")
            .execute()
            .data
            or []
        )

        return {
            "people_total": len(people),
            "women_total": sum(
                1 for r in people if r.get("entity_type") == "woman"
            ),
            "men_total": sum(
                1 for r in people if r.get("entity_type") == "man"
            ),
            "high_priority": sum(
                1 for r in people if r.get("priority") == "High"
            ),
            "series_total": len(series),
        }

    con = _local_connect()
    out = {
        "people_total": con.execute(
            "SELECT COUNT(*) FROM catalog_people"
        ).fetchone()[0],
        "women_total": con.execute(
            "SELECT COUNT(*) FROM catalog_people WHERE entity_type='woman'"
        ).fetchone()[0],
        "men_total": con.execute(
            "SELECT COUNT(*) FROM catalog_people WHERE entity_type='man'"
        ).fetchone()[0],
        "high_priority": con.execute(
            "SELECT COUNT(*) FROM catalog_people WHERE priority='High'"
        ).fetchone()[0],
        "series_total": con.execute(
            "SELECT COUNT(*) FROM catalog_series"
        ).fetchone()[0],
    }
    con.close()
    return out


def list_catalog_people(
    search="",
    priority="",
    entity_type="",
    playlist="",
    limit=500,
):
    if using_supabase():
        query = _client().table("catalog_people").select("*")

        if search:
            safe = search.replace("%", "")
            query = query.or_(
                f"name.ilike.%{safe}%,"
                f"story_role.ilike.%{safe}%,"
                f"bible_references.ilike.%{safe}%"
            )

        if priority:
            query = query.eq("priority", priority)

        if entity_type:
            query = query.eq("entity_type", entity_type)

        if playlist:
            query = query.eq("playlist_series", playlist)

        rows = query.limit(int(limit)).execute().data or []

        priority_rank = {
            "High": 1,
            "Medium": 2,
            "Low": 3,
        }

        rows.sort(
            key=lambda r: (
                priority_rank.get(r.get("priority"), 4),
                1 if r.get("cinematic_score") is None else 0,
                -(r.get("cinematic_score") or 0),
                r.get("source_position") or 999999,
                (r.get("name") or "").lower(),
            )
        )
        return rows

    clauses = []
    params = []

    if search:
        clauses.append(
            "(name LIKE ? OR story_role LIKE ? OR bible_references LIKE ?)"
        )
        q = f"%{search}%"
        params += [q, q, q]

    if priority:
        clauses.append("priority=?")
        params.append(priority)

    if entity_type:
        clauses.append("entity_type=?")
        params.append(entity_type)

    if playlist:
        clauses.append("playlist_series=?")
        params.append(playlist)

    where = (" WHERE " + " AND ".join(clauses)) if clauses else ""

    sql = f"""
    SELECT * FROM catalog_people {where}
    ORDER BY
      CASE priority
        WHEN 'High' THEN 1
        WHEN 'Medium' THEN 2
        WHEN 'Low' THEN 3
        ELSE 4
      END,
      CASE WHEN cinematic_score IS NULL THEN 1 ELSE 0 END,
      cinematic_score DESC,
      source_position ASC,
      name ASC
    LIMIT ?
    """

    params.append(int(limit))

    con = _local_connect()
    rows = [dict(r) for r in con.execute(sql, params)]
    con.close()
    return rows


def list_series():
    if using_supabase():
        result = (
            _client()
            .table("catalog_series")
            .select("*")
            .order("rank")
            .execute()
        )
        return result.data or []

    con = _local_connect()
    rows = [
        dict(r)
        for r in con.execute(
            """
            SELECT * FROM catalog_series
            ORDER BY rank ASC, series_name ASC
            """
        )
    ]
    con.close()
    return rows


def get_catalog_person(person_id):
    if using_supabase():
        result = (
            _client()
            .table("catalog_people")
            .select("*")
            .eq("id", person_id)
            .limit(1)
            .execute()
        )
        return result.data[0] if result.data else None

    con = _local_connect()
    row = con.execute(
        "SELECT * FROM catalog_people WHERE id=?",
        (person_id,),
    ).fetchone()
    con.close()
    return _local_row(row)


def set_catalog_production_status(person_id, status):
    allowed = {"not_started", "in_progress", "completed"}
    if status not in allowed:
        raise ValueError("Invalid production status.")

    if using_supabase():
        payload = {
            "production_status": status,
            "completed_at": None,
        }
        if status == "completed":
            from datetime import datetime, timezone
            payload["completed_at"] = datetime.now(timezone.utc).isoformat()

        (
            _client()
            .table("catalog_people")
            .update(payload)
            .eq("id", person_id)
            .execute()
        )
        return

    con = _local_connect()
    con.execute(
        """
        UPDATE catalog_people
        SET production_status=?,
            completed_at=CASE
              WHEN ?='completed' THEN CURRENT_TIMESTAMP
              ELSE NULL
            END,
            updated_at=CURRENT_TIMESTAMP
        WHERE id=?
        """,
        (status, status, person_id),
    )
    con.commit()
    con.close()


def catalog_progress_summary():
    rows = list_catalog_people(limit=5000)
    total = len(rows)
    completed = sum(
        1 for row in rows
        if row.get("production_status") == "completed"
    )
    in_progress = sum(
        1 for row in rows
        if row.get("production_status") == "in_progress"
    )
    not_started = max(total - completed - in_progress, 0)

    return {
        "total": total,
        "completed": completed,
        "in_progress": in_progress,
        "not_started": not_started,
    }


def link_project(
    project_id,
    person_id=None,
    series_id=None,
    episode_title=None,
):
    payload = {
        "project_id": project_id,
        "catalog_person_id": person_id,
        "catalog_series_id": series_id,
        "episode_title": episode_title,
    }

    if using_supabase():
        (
            _client()
            .table("project_catalog_links")
            .upsert(payload, on_conflict="project_id")
            .execute()
        )
        return

    con = _local_connect()
    con.execute(
        """
        INSERT INTO project_catalog_links(
            project_id,
            catalog_person_id,
            catalog_series_id,
            episode_title
        )
        VALUES(?,?,?,?)
        ON CONFLICT(project_id) DO UPDATE SET
          catalog_person_id=excluded.catalog_person_id,
          catalog_series_id=excluded.catalog_series_id,
          episode_title=excluded.episode_title
        """,
        (
            project_id,
            person_id,
            series_id,
            episode_title,
        ),
    )
    con.commit()
    con.close()


def _stable(prefix, *parts):
    raw = "|".join(
        str(part or "").strip().lower()
        for part in parts
    )
    return prefix + "_" + hashlib.sha1(raw.encode()).hexdigest()[:12]


def _sheet_rows(workbook, name):
    ws = workbook[name]
    headers = [
        str(cell.value or "").strip()
        for cell in ws[1]
    ]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        yield dict(zip(headers, row))


def _first(row, *keys):
    for key in keys:
        value = row.get(key)
        if value not in (None, ""):
            return value
    return None


def _clean(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    return str(value).strip()


def _float_or_none(value):
    try:
        if value in (None, ""):
            return None
        return float(value)
    except Exception:
        return None


def _int_or_none(value):
    try:
        if value in (None, ""):
            return None
        return int(value)
    except Exception:
        return None


def seed_catalog_if_empty():
    """
    Populate the cloud catalogue from the bundled planning seed exactly once.
    This removes the need for a user to upload the spreadsheet during normal use.
    """
    summary = catalog_summary()
    if summary["people_total"] > 0 or summary["series_total"] > 0:
        return summary

    seed_path = Path(__file__).with_name("catalog_seed.json")
    if not seed_path.exists():
        return summary

    data = json.loads(seed_path.read_text(encoding="utf-8"))

    people_rows = []
    series_rows = []

    for row in data.get("master_checklist", []):
        name = _first(row, "Woman / Character", "Character", "Name")
        if not name:
            continue

        references = _first(
            row,
            "Bible Reference(s)",
            "Bible References",
            "Bible Reference",
        )

        people_rows.append(
            {
                "id": _stable(
                    "person",
                    "Master Checklist",
                    name,
                    references,
                ),
                "source_sheet": "Master Checklist",
                "source_position": _int_or_none(
                    _first(row, "Order", "Rank")
                ),
                "testament": _clean(_first(row, "Testament")),
                "book": _clean(_first(row, "Book")),
                "name": str(name).strip(),
                "entity_type": "woman",
                "named_status": _clean(
                    _first(row, "Named Status")
                ),
                "bible_references": _clean(references),
                "story_role": _clean(
                    _first(
                        row,
                        "Story / Role",
                        "Story/Role",
                        "Story Role",
                    )
                ),
                "priority": _clean(_first(row, "Priority")),
                "cinematic_score": _float_or_none(
                    _first(row, "Cinematic Score")
                ),
                "youtube_hook": _clean(
                    _first(row, "YouTube Hook")
                ),
                "reference_asset_need": _clean(
                    _first(row, "Reference Asset Need")
                ),
                "primary_environment": _clean(
                    _first(row, "Primary Environment")
                ),
                "canon": _clean(_first(row, "Canon")),
                "playlist_series": _clean(
                    _first(
                        row,
                        "Playlist / Series",
                        "Playlist/Series",
                    )
                ),
                "notes": _clean(_first(row, "Notes")),
                "source_url": _clean(
                    _first(row, "Source URL")
                ),
                "ai_studio_audio_status": _clean(
                    _first(row, "AI Studio Audio")
                ),
                "flow_reference_asset_status": _clean(
                    _first(row, "Flow Reference Asset")
                ),
                "flow_scene_prompt_status": _clean(
                    _first(row, "Flow Scene Prompt")
                ),
                "animation_status": _clean(
                    _first(row, "Grok Animation", "Animation")
                ),
                "final_edit_status": _clean(
                    _first(row, "Final Edit")
                ),
                "ready_status": _clean(
                    _first(row, "Ready?", "Ready")
                ),
            }
        )

    for row in data.get("men_ranked", []):
        name = _first(row, "Man / Character", "Character", "Name")
        if not name:
            continue

        references = _first(
            row,
            "Bible Reference(s)",
            "Bible References",
            "Bible Reference",
        )

        people_rows.append(
            {
                "id": _stable(
                    "person",
                    "Men Ranked",
                    name,
                    references,
                ),
                "source_sheet": "Men Ranked",
                "source_position": _int_or_none(
                    _first(row, "Rank", "Order")
                ),
                "testament": _clean(_first(row, "Testament")),
                "book": _clean(_first(row, "Book")),
                "name": str(name).strip(),
                "entity_type": "man",
                "named_status": _clean(
                    _first(row, "Named Status")
                ),
                "bible_references": _clean(references),
                "story_role": _clean(
                    _first(
                        row,
                        "Story / Role",
                        "Story/Role",
                        "Story Role",
                    )
                ),
                "priority": _clean(_first(row, "Priority")),
                "cinematic_score": _float_or_none(
                    _first(row, "Cinematic Score")
                ),
                "youtube_hook": _clean(
                    _first(row, "YouTube Hook")
                ),
                "reference_asset_need": _clean(
                    _first(row, "Reference Asset Need")
                ),
                "primary_environment": _clean(
                    _first(row, "Primary Environment")
                ),
                "canon": _clean(_first(row, "Canon")),
                "playlist_series": _clean(
                    _first(
                        row,
                        "Playlist / Series",
                        "Playlist/Series",
                    )
                ),
                "notes": _clean(_first(row, "Notes")),
                "source_url": _clean(
                    _first(row, "Source URL")
                ),
                "ai_studio_audio_status": _clean(
                    _first(row, "AI Studio Audio")
                ),
                "flow_reference_asset_status": _clean(
                    _first(row, "Flow Reference Asset")
                ),
                "flow_scene_prompt_status": _clean(
                    _first(row, "Flow Scene Prompt")
                ),
                "animation_status": _clean(
                    _first(row, "Grok Animation", "Animation")
                ),
                "final_edit_status": _clean(
                    _first(row, "Final Edit")
                ),
                "ready_status": _clean(
                    _first(row, "Ready?", "Ready")
                ),
            }
        )

    for row in data.get("series_playlists", []):
        name = _first(
            row,
            "Series / Playlist",
            "Series/Playlist",
            "Series",
        )
        if not name:
            continue

        series_rows.append(
            {
                "id": _stable("series", name),
                "rank": _int_or_none(_first(row, "Rank")),
                "series_name": str(name).strip(),
                "series_type": _clean(_first(row, "Type")),
                "core_concept": _clean(
                    _first(row, "Core Concept")
                ),
                "best_starter_episodes": _clean(
                    _first(
                        row,
                        "Best Starter Episodes",
                        "Starter Episodes",
                    )
                ),
                "youtube_reason": _clean(
                    _first(
                        row,
                        "Why It Can Work on YouTube",
                        "YouTube Reason",
                    )
                ),
                "recommended_scene_count": _clean(
                    _first(
                        row,
                        "Recommended Scene Count",
                        "Scene Count",
                    )
                ),
                "recommended_runtime": _clean(
                    _first(
                        row,
                        "Recommended Runtime",
                        "Runtime",
                    )
                ),
                "priority": _clean(_first(row, "Priority")),
                "status": _clean(_first(row, "Status")),
                "source_notes": _clean(
                    _first(
                        row,
                        "Source / Notes",
                        "Source Notes",
                    )
                ),
            }
        )

    if using_supabase():
        client = _client()
        chunk_size = 100

        for start in range(0, len(people_rows), chunk_size):
            (
                client.table("catalog_people")
                .upsert(
                    people_rows[start:start + chunk_size],
                    on_conflict="id",
                )
                .execute()
            )

        for start in range(0, len(series_rows), chunk_size):
            (
                client.table("catalog_series")
                .upsert(
                    series_rows[start:start + chunk_size],
                    on_conflict="id",
                )
                .execute()
            )

        return catalog_summary()

    con = _local_connect()

    for row in people_rows:
        cols = list(row.keys())
        values = [row[col] for col in cols]
        placeholders = ",".join("?" for _ in cols)
        con.execute(
            f"INSERT OR REPLACE INTO catalog_people({','.join(cols)}) VALUES({placeholders})",
            values,
        )

    for row in series_rows:
        cols = list(row.keys())
        values = [row[col] for col in cols]
        placeholders = ",".join("?" for _ in cols)
        con.execute(
            f"INSERT OR REPLACE INTO catalog_series({','.join(cols)}) VALUES({placeholders})",
            values,
        )

    con.commit()
    con.close()
    return catalog_summary()


def import_spreadsheet(file_bytes):
    workbook = load_workbook(
        BytesIO(file_bytes),
        data_only=True,
    )

    people_rows = []
    series_rows = []

    if "Master Checklist" in workbook.sheetnames:
        for row in _sheet_rows(workbook, "Master Checklist"):
            name = _first(
                row,
                "Woman / Character",
                "Character",
                "Name",
            )

            if not name:
                continue

            references = _first(
                row,
                "Bible Reference(s)",
                "Bible References",
                "Bible Reference",
            )

            people_rows.append(
                {
                    "id": _stable(
                        "person",
                        "Master Checklist",
                        name,
                        references,
                    ),
                    "source_sheet": "Master Checklist",
                    "source_position": _int_or_none(
                        _first(row, "Order", "Rank")
                    ),
                    "testament": _clean(
                        _first(row, "Testament")
                    ),
                    "book": _clean(_first(row, "Book")),
                    "name": str(name).strip(),
                    "entity_type": "woman",
                    "named_status": _clean(
                        _first(row, "Named Status")
                    ),
                    "bible_references": _clean(references),
                    "story_role": _clean(
                        _first(
                            row,
                            "Story / Role",
                            "Story/Role",
                            "Story Role",
                        )
                    ),
                    "priority": _clean(
                        _first(row, "Priority")
                    ),
                    "cinematic_score": _float_or_none(
                        _first(row, "Cinematic Score")
                    ),
                    "youtube_hook": _clean(
                        _first(row, "YouTube Hook")
                    ),
                    "reference_asset_need": _clean(
                        _first(row, "Reference Asset Need")
                    ),
                    "primary_environment": _clean(
                        _first(row, "Primary Environment")
                    ),
                    "canon": _clean(_first(row, "Canon")),
                    "playlist_series": _clean(
                        _first(
                            row,
                            "Playlist / Series",
                            "Playlist/Series",
                        )
                    ),
                    "notes": _clean(_first(row, "Notes")),
                    "source_url": _clean(
                        _first(row, "Source URL")
                    ),
                }
            )

    if "Men Ranked" in workbook.sheetnames:
        for row in _sheet_rows(workbook, "Men Ranked"):
            name = _first(
                row,
                "Man / Character",
                "Character",
                "Name",
            )

            if not name:
                continue

            references = _first(
                row,
                "Bible Reference(s)",
                "Bible References",
                "Bible Reference",
            )

            people_rows.append(
                {
                    "id": _stable(
                        "person",
                        "Men Ranked",
                        name,
                        references,
                    ),
                    "source_sheet": "Men Ranked",
                    "source_position": _int_or_none(
                        _first(row, "Rank", "Order")
                    ),
                    "testament": _clean(
                        _first(row, "Testament")
                    ),
                    "book": _clean(_first(row, "Book")),
                    "name": str(name).strip(),
                    "entity_type": "man",
                    "named_status": _clean(
                        _first(row, "Named Status")
                    ),
                    "bible_references": _clean(references),
                    "story_role": _clean(
                        _first(
                            row,
                            "Story / Role",
                            "Story/Role",
                            "Story Role",
                        )
                    ),
                    "priority": _clean(
                        _first(row, "Priority")
                    ),
                    "cinematic_score": _float_or_none(
                        _first(row, "Cinematic Score")
                    ),
                    "youtube_hook": _clean(
                        _first(row, "YouTube Hook")
                    ),
                    "reference_asset_need": _clean(
                        _first(row, "Reference Asset Need")
                    ),
                    "primary_environment": _clean(
                        _first(row, "Primary Environment")
                    ),
                    "canon": _clean(_first(row, "Canon")),
                    "playlist_series": _clean(
                        _first(
                            row,
                            "Playlist / Series",
                            "Playlist/Series",
                        )
                    ),
                    "notes": _clean(_first(row, "Notes")),
                    "source_url": _clean(
                        _first(row, "Source URL")
                    ),
                }
            )

    if "Series Playlists" in workbook.sheetnames:
        for row in _sheet_rows(workbook, "Series Playlists"):
            name = _first(
                row,
                "Series / Playlist",
                "Series/Playlist",
                "Series",
            )

            if not name:
                continue

            series_rows.append(
                {
                    "id": _stable("series", name),
                    "rank": _int_or_none(
                        _first(row, "Rank")
                    ),
                    "series_name": str(name).strip(),
                    "series_type": _clean(
                        _first(row, "Type")
                    ),
                    "core_concept": _clean(
                        _first(row, "Core Concept")
                    ),
                    "best_starter_episodes": _clean(
                        _first(
                            row,
                            "Best Starter Episodes",
                            "Starter Episodes",
                        )
                    ),
                    "youtube_reason": _clean(
                        _first(
                            row,
                            "Why It Can Work on YouTube",
                            "YouTube Reason",
                        )
                    ),
                    "recommended_scene_count": _clean(
                        _first(
                            row,
                            "Recommended Scene Count",
                            "Scene Count",
                        )
                    ),
                    "recommended_runtime": _clean(
                        _first(
                            row,
                            "Recommended Runtime",
                            "Runtime",
                        )
                    ),
                    "priority": _clean(
                        _first(row, "Priority")
                    ),
                    "status": _clean(
                        _first(row, "Status")
                    ),
                    "source_notes": _clean(
                        _first(
                            row,
                            "Source / Notes",
                            "Source Notes",
                        )
                    ),
                }
            )

    if using_supabase():
        client = _client()

        chunk_size = 100

        for start in range(0, len(people_rows), chunk_size):
            client.table("catalog_people").upsert(
                people_rows[start:start + chunk_size],
                on_conflict="id",
            ).execute()

        for start in range(0, len(series_rows), chunk_size):
            client.table("catalog_series").upsert(
                series_rows[start:start + chunk_size],
                on_conflict="id",
            ).execute()

        return catalog_summary()

    con = _local_connect()

    for row in people_rows:
        cols = list(row.keys())
        values = [row[col] for col in cols]
        placeholders = ",".join("?" for _ in cols)

        con.execute(
            f"""
            INSERT OR REPLACE INTO catalog_people(
                {",".join(cols)}
            )
            VALUES({placeholders})
            """,
            values,
        )

    for row in series_rows:
        cols = list(row.keys())
        values = [row[col] for col in cols]
        placeholders = ",".join("?" for _ in cols)

        con.execute(
            f"""
            INSERT OR REPLACE INTO catalog_series(
                {",".join(cols)}
            )
            VALUES({placeholders})
            """,
            values,
        )

    con.commit()
    con.close()

    return catalog_summary()
